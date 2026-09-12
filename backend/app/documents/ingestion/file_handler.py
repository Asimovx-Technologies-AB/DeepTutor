"""
File Handler — File format detection, validation, and multi-format opening.

Consolidates file-type dispatch logic from the former doc_processor.py and
study_doc_processor.py into a single, reusable component.
"""
from __future__ import annotations

import asyncio
import hashlib
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.documents.models import PipelineDocument, PipelineStatus

logger = logging.getLogger(__name__)


# ─── Format registry ─────────────────────────────────────────────────────────

FORMAT_MAP: Dict[str, set] = {
    "pdf": {".pdf"},
    "docx": {".docx", ".doc"},
    "pptx": {".pptx", ".ppt"},
    "image": {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"},
    "text": {".txt", ".md", ".py", ".csv", ".json"},
}

ALL_SUPPORTED_EXTENSIONS = set()
for _exts in FORMAT_MAP.values():
    ALL_SUPPORTED_EXTENSIONS |= _exts


class FileHandler:
    """Detects file format, validates constraints, and opens files."""

    # ── Format detection ──────────────────────────────────────────────────

    @staticmethod
    def detect_format(file_path: str) -> str:
        """Returns normalised format key: 'pdf', 'docx', 'pptx', 'image', 'text'.

        Raises ValueError for unsupported formats.
        """
        ext = Path(file_path).suffix.lower()
        for fmt, extensions in FORMAT_MAP.items():
            if ext in extensions:
                return fmt
        raise ValueError(
            f"Unsupported file format '{ext}'. "
            f"Supported: {sorted(ALL_SUPPORTED_EXTENSIONS)}"
        )

    # ── Validation ────────────────────────────────────────────────────────

    @staticmethod
    def validate(file_path: str, max_size_mb: int = 200) -> Tuple[bool, Optional[str]]:
        """Validates file existence, size, and format support.

        Returns (True, None) on success or (False, error_message) on failure.
        """
        path = Path(file_path)
        if not path.exists():
            return False, f"File not found: {file_path}"
        if not path.is_file():
            return False, f"Path is not a file: {file_path}"

        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > max_size_mb:
            return False, f"File too large ({size_mb:.1f} MB > {max_size_mb} MB limit)"

        ext = path.suffix.lower()
        if ext not in ALL_SUPPORTED_EXTENSIONS:
            return False, (
                f"Unsupported format '{ext}'. "
                f"Supported: {sorted(ALL_SUPPORTED_EXTENSIONS)}"
            )

        return True, None

    # ── Content hashing ───────────────────────────────────────────────────

    @staticmethod
    def compute_hash(file_path: str) -> str:
        """SHA-256 content hash for deduplication."""
        hasher = hashlib.sha256()
        with open(file_path, "rb") as f:
            for block in iter(lambda: f.read(65536), b""):
                hasher.update(block)
        return hasher.hexdigest()

    # ── Format-specific openers ───────────────────────────────────────────

    async def open_pdf(self, file_path: str) -> Tuple[int, Any]:
        """Opens a PDF with PyMuPDF and returns (page_count, doc_handle).

        The caller is responsible for closing the document handle.
        """
        import pymupdf

        def _open():
            doc = pymupdf.open(file_path)
            return len(doc), doc

        return await asyncio.to_thread(_open)

    async def open_docx(self, file_path: str) -> Tuple[int, List[Tuple[int, str]]]:
        """Extracts paragraphs + table rows from a Word document.

        Returns (page_count=1, [(paragraph_idx, text)]).
        """
        def _extract():
            import docx

            doc = docx.Document(file_path)
            parts: List[Tuple[int, str]] = []
            idx = 0
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    parts.append((idx, txt))
                    idx += 1
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        parts.append((idx, row_text))
                        idx += 1
            return parts

        parts = await asyncio.to_thread(_extract)
        return 1, parts

    async def open_pptx(self, file_path: str) -> Tuple[int, List[Tuple[int, str]]]:
        """Extracts per-slide text from a PowerPoint presentation.

        Returns (slide_count, [(slide_num_1based, slide_text)]).
        """
        def _extract():
            from pptx import Presentation

            prs = Presentation(file_path)
            slides: List[Tuple[int, str]] = []
            for idx, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
                if texts:
                    slides.append((idx + 1, "\n".join(texts)))
            return len(prs.slides), slides

        count, slides = await asyncio.to_thread(_extract)
        return count, slides

    async def read_text(self, file_path: str) -> str:
        """Reads a plain text / code / markdown file."""
        def _read():
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        return await asyncio.to_thread(_read)

    async def read_image(self, file_path: str) -> bytes:
        """Reads an image file as raw bytes."""
        def _read():
            with open(file_path, "rb") as f:
                return f.read()

        return await asyncio.to_thread(_read)

    # ── Convenience: create PipelineDocument ──────────────────────────────

    def create_pipeline_document(
        self,
        doc_id: str,
        file_path: str,
        file_name: str,
        session_id: str = "",
        user_id: Optional[str] = None,
        subject: str = "General",
        doc_hash: Optional[str] = None,
    ) -> PipelineDocument:
        """Creates a PipelineDocument with format detected and hash computed."""
        file_type = self.detect_format(file_path)
        if doc_hash is None:
            doc_hash = self.compute_hash(file_path)
        return PipelineDocument(
            doc_id=doc_id,
            file_path=file_path,
            file_name=file_name,
            file_type=file_type,
            session_id=session_id,
            user_id=user_id,
            subject=subject,
            status=PipelineStatus.RECEIVED,
            doc_hash=doc_hash,
        )
